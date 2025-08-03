import os
import PyPDF2
import docx
from typing import Optional, Tuple
import logging

logger = logging.getLogger(__name__)

def extract_text_from_file(file_path: str) -> Tuple[str, Optional[str]]:
    """
    Extract text from uploaded files (PDF, DOCX, TXT)
    Returns: (extracted_text, error_message)
    """
    try:
        file_extension = os.path.splitext(file_path)[1].lower()
        
        if file_extension == '.pdf':
            return extract_pdf_text(file_path)
        elif file_extension in ['.docx', '.doc']:
            return extract_docx_text(file_path)
        elif file_extension == '.txt':
            return extract_txt_text(file_path)
        else:
            return "", f"Unsupported file format: {file_extension}"
            
    except Exception as e:
        logger.error(f"Error extracting text from {file_path}: {str(e)}")
        return "", f"Error processing file: {str(e)}"

def extract_pdf_text(file_path: str) -> Tuple[str, Optional[str]]:
    """Extract text from PDF file"""
    try:
        text = ""
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
        
        if not text.strip():
            return "", "No text could be extracted from PDF"
        
        return text.strip(), None
        
    except Exception as e:
        return "", f"Error reading PDF: {str(e)}"

def extract_docx_text(file_path: str) -> Tuple[str, Optional[str]]:
    """Extract text from DOCX file"""
    try:
        doc = docx.Document(file_path)
        text = ""
        
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        
        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
        
        if not text.strip():
            return "", "No text could be extracted from DOCX"
        
        return text.strip(), None
        
    except Exception as e:
        return "", f"Error reading DOCX: {str(e)}"

def extract_txt_text(file_path: str) -> Tuple[str, Optional[str]]:
    """Extract text from TXT file"""
    try:
        encodings = ['utf-8', 'latin-1', 'cp1252']
        
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as file:
                    text = file.read()
                    return text.strip(), None
            except UnicodeDecodeError:
                continue
        
        return "", "Could not decode text file with any supported encoding"
        
    except Exception as e:
        return "", f"Error reading TXT: {str(e)}"

def clean_extracted_text(text: str) -> str:
    """Clean and normalize extracted text"""
    import re
    
    # Remove excessive whitespace
    text = re.sub(r'\s+', ' ', text)
    
    # Remove special characters but keep punctuation
    text = re.sub(r'[^\w\s\.\,\;\:\!\?\-\(\)\[\]\'\"]+', ' ', text)
    
    # Remove extra spaces
    text = ' '.join(text.split())
    
    return text.strip()

def estimate_reading_time(text: str, words_per_minute: int = 200) -> float:
    """Estimate reading time in hours based on word count"""
    word_count = len(text.split())
    reading_time_minutes = word_count / words_per_minute
    return reading_time_minutes / 60  # Convert to hours

def calculate_text_complexity(text: str) -> float:
    """
    Calculate text complexity score (0-1 scale)
    Based on sentence length, word length, and vocabulary diversity
    """
    import string
    
    if not text:
        return 0.0
    
    # Basic metrics
    sentences = text.count('.') + text.count('!') + text.count('?')
    words = text.split()
    word_count = len(words)
    
    if sentences == 0 or word_count == 0:
        return 0.0
    
    # Average sentence length
    avg_sentence_length = word_count / sentences
    
    # Average word length
    clean_words = [word.strip(string.punctuation) for word in words if word.strip(string.punctuation)]
    avg_word_length = sum(len(word) for word in clean_words) / len(clean_words) if clean_words else 0
    
    # Vocabulary diversity (unique words / total words)
    unique_words = len(set(word.lower() for word in clean_words))
    vocab_diversity = unique_words / len(clean_words) if clean_words else 0
    
    # Normalize metrics and combine
    sentence_complexity = min(avg_sentence_length / 20, 1.0)  # Max 20 words per sentence
    word_complexity = min(avg_word_length / 8, 1.0)  # Max 8 characters per word
    vocab_complexity = min(vocab_diversity, 1.0)
    
    # Weighted combination
    complexity_score = (sentence_complexity * 0.4 + word_complexity * 0.3 + vocab_complexity * 0.3)
    
    return min(complexity_score, 1.0)
